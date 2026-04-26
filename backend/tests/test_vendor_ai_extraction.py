"""Tests for Vendor Audit AI extraction (Phase 1 pivot).

Live Gemini Flash + Sonnet 4.5 calls, but minimal — we mostly verify the routing,
validation, caching, and combiner glue.
"""
import io
import os
import time
import pytest
import requests
import openpyxl

BASE_URL = ""
with open("/app/frontend/.env") as f:
    for line in f:
        if line.startswith("REACT_APP_BACKEND_URL="):
            BASE_URL = line.split("=", 1)[1].strip().rstrip("/")


def _login(email, password):
    r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_token():
    return _login("admin@hrms.com", "admin123")


def _h(t):
    return {"Authorization": f"Bearer {t}"}


def _wage_register_xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wage Register"
    ws["A1"] = "Wage Register — March 2026"
    ws["A2"] = "Establishment: Acme Manufacturing Pvt Ltd, Pune"
    headers = ["UAN", "Employee Name", "Designation", "Days Worked", "Basic", "DA",
               "Gross", "EPF Wages", "EPF Contribution", "ESIC Wages", "ESIC Employee", "PT", "Net Pay"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i).value = h
    rows = [
        ("100123456789", "Rahul S", "Operator", 26, 15000, 5000, 20000, 15000, 1800, 20000, 150, 200, 17850),
        ("100123456790", "Priya P", "Supervisor", 26, 22000, 6000, 28000, 15000, 1800, 21000, 158, 200, 25842),
    ]
    for ri, row in enumerate(rows, 5):
        for ci, v in enumerate(row, 1):
            ws.cell(row=ri, column=ci).value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture(scope="module")
def audit_id(admin_token):
    """Create a fresh audit for these tests."""
    h = _h(admin_token)
    contractor = requests.get(f"{BASE_URL}/api/vendor-audit/contractors", headers=h).json()[0]
    r = requests.post(
        f"{BASE_URL}/api/vendor-audit/audits/start?contractor_id={contractor['id']}",
        headers=h, json={"wage_month": "JUL-2026", "state": "MAHARASHTRA"},
    )
    assert r.status_code == 200, r.text
    aid = r.json()["id"]
    yield aid
    requests.delete(f"{BASE_URL}/api/vendor-audit/audits/{aid}", headers=h)


def test_meta_doc_types(admin_token):
    r = requests.get(f"{BASE_URL}/api/vendor-audit/meta/doc-types", headers=_h(admin_token))
    assert r.status_code == 200
    items = r.json()
    assert any(x["key"] == "pf_ecr" for x in items)
    assert any(x["key"] == "esic_paid_challan" for x in items)


def test_upload_invalid_filetype(admin_token, audit_id):
    files = {"file": ("notes.txt", b"hi", "text/plain")}
    r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/upload-document",
                       headers=_h(admin_token), files=files)
    assert r.status_code == 400


def test_upload_xlsx_extracts_employees_live(admin_token, audit_id):
    """Live AI test: upload a wage register Excel, confirm AI extracts employee rows."""
    files = {"file": ("wage_reg.xlsx", _wage_register_xlsx(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/upload-document",
                       headers=_h(admin_token), files=files)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["sha256"]
    assert "extracted" in d
    assert d["validation"]["status"] in ("valid", "mismatch", "unknown")
    employees = (d.get("extracted") or {}).get("employees") or []
    assert len(employees) >= 1
    assert d["cost_inr"] > 0  # paid live (not cached)


def test_upload_same_file_hits_cache(admin_token, audit_id):
    """Re-upload the SAME file → from_cache=True, cost=0."""
    payload = _wage_register_xlsx()  # generated once — same bytes both calls
    files1 = {"file": ("wage_reg2.xlsx", payload,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r1 = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/upload-document",
                       headers=_h(admin_token), files=files1)
    assert r1.status_code == 200
    files2 = {"file": ("wage_reg2_dupe.xlsx", payload,
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r2 = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/upload-document",
                       headers=_h(admin_token), files=files2)
    assert r2.status_code == 200
    d = r2.json()
    assert d["from_cache"] is True
    assert d["cost_inr"] == 0


def test_list_and_delete_ai_documents(admin_token, audit_id):
    h = _h(admin_token)
    r = requests.get(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/ai-documents", headers=h)
    assert r.status_code == 200
    items = r.json()
    assert len(items) >= 1
    # Ensure file_b64 is stripped from list response
    assert all("file_b64" not in it for it in items)
    # Delete
    doc_id = items[0]["id"]
    rd = requests.delete(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/ai-documents/{doc_id}", headers=h)
    assert rd.status_code == 200


def test_run_ai_audit_with_extracted_data(admin_token, audit_id):
    """End-to-end: upload one doc, run-ai → rows populated, findings produced."""
    h = _h(admin_token)
    files = {"file": ("wage_reg.xlsx", _wage_register_xlsx(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/upload-document",
                   headers=h, files=files)
    r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/run-ai", headers=h)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["ai_doc_count"] >= 1
    assert d["ai_rows_count"] >= 1
    assert "totals" in d
    # Should have flagged missing PF challan / ESIC
    assert any("Missing document" in (f.get("title") or "") for f in d.get("summary_findings") or [])


def test_run_ai_without_docs_400(admin_token, audit_id):
    h = _h(admin_token)
    # Delete every existing AI doc first
    items = requests.get(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/ai-documents", headers=h).json()
    for it in items:
        requests.delete(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/ai-documents/{it['id']}", headers=h)
    r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{audit_id}/run-ai", headers=h)
    assert r.status_code == 400


def test_validators_unit():
    """Direct unit test on the regex rules."""
    import sys
    sys.path.insert(0, "/app/backend")
    from vendor_audit.validators import validate_document_text
    pf_text = "Combined Challan A/C No. 01, 02, 10, 21 & 22\nProvident Fund Organisation\nTRRN: 123\nEstablishment Code: MH-PUN-12345\nWage Month: MAR-2026"
    status, conf, detected, msg = validate_document_text(pf_text, "pf_challan")
    assert status == "valid"
    assert detected == "pf_challan"
    # mismatch when uploaded under wrong type
    status2, _, det2, _ = validate_document_text(pf_text, "esic_paid_challan")
    assert status2 == "mismatch"
    assert det2 == "pf_challan"
    # unknown for unrelated text
    status3, _, _, _ = validate_document_text("This is just a leave letter, not a statutory doc.", "")
    assert status3 == "unknown"


def test_combiner_unit():
    import sys
    sys.path.insert(0, "/app/backend")
    from vendor_audit.combiner import combine_extractions
    docs = [
        {
            "claimed_doc_type": "wage_register",
            "extracted": {
                "doc_type_detected": "wage_register",
                "summary": {"establishment_name": "Acme"},
                "employees": [
                    {"uan": "100", "name": "Rahul", "gross_wages": 20000, "epf_wages": 15000, "epf_contribution": 1800},
                    {"uan": "200", "name": "Priya", "gross_wages": 28000, "epf_wages": 15000, "epf_contribution": 1800},
                ], "confidence": 0.9,
            },
            "validation": {"status": "valid", "detected_type": "wage_register"},
        },
    ]
    rows, parsed, warnings = combine_extractions(docs)
    assert len(rows) == 2
    assert rows[0]["UAN"] == "100"
    assert rows[0]["FINAL GROSS EARNED"] == 20000
    assert "wage_register" in parsed


def test_extraction_cache_endpoint(admin_token):
    r = requests.get(f"{BASE_URL}/api/vendor-audit/cache/stats", headers=_h(admin_token))
    assert r.status_code == 200
    assert "cached_extractions" in r.json()

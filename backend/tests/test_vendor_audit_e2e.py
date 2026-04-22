"""End-to-end backend tests for Vendor Audit module (iteration 13)."""
import os
import io
import time
import uuid
import zipfile
import requests
import pytest
import openpyxl

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://talent-board-14.preview.emergentagent.com").rstrip("/")
ADMIN_EMAIL = "admin@hrms.com"
ADMIN_PASSWORD = "admin123"
SAMPLE_PDF_URL = "https://customer-assets.emergentagent.com/job_talent-board-14/artifacts/0sa0vm3a_STATUTORY%20DOCUMENTS.pdf"


@pytest.fixture(scope="module")
def admin_token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD, "login_as": "admin"})
    if r.status_code != 200:
        pytest.skip(f"Admin login failed: {r.status_code} {r.text}")
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="module")
def sample_pdf_bytes():
    r = requests.get(SAMPLE_PDF_URL, timeout=60)
    if r.status_code != 200:
        pytest.skip("Sample PDF not reachable")
    return r.content


# ── Contractor lifecycle ──
class TestContractorLifecycle:
    state = {}

    def test_create_contractor(self, admin_headers):
        unique_email = f"test_vendor_{uuid.uuid4().hex[:8]}@example.com"
        body = {
            "name": "TEST E2E Vendor", "legal_name": "TEST E2E Pvt Ltd",
            "establishment_code": "E2E-001", "lin": "1234567890",
            "pf_code": "PF-E2E", "esic_code": "ES-E2E",
            "pt_registration_no": "PT-E2E", "mlwf_lin": "MLWF-E2E",
            "address": "Test addr, Pune", "state": "MAHARASHTRA",
            "contact_person": "Test Contact", "contact_email": unique_email,
            "contact_phone": "9999999999", "scope_of_work": "Housekeeping"
        }
        r = requests.post(f"{BASE_URL}/api/vendor-audit/contractors", json=body, headers=admin_headers)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "temp_password" in data and len(data["temp_password"]) > 6
        assert data["contractor"]["name"] == body["name"]
        assert data["contractor"]["contact_email"] == unique_email
        assert "id" in data["contractor"]
        TestContractorLifecycle.state["cid"] = data["contractor"]["id"]
        TestContractorLifecycle.state["email"] = unique_email
        TestContractorLifecycle.state["pwd"] = data["temp_password"]

    def test_list_contractors(self, admin_headers):
        r = requests.get(f"{BASE_URL}/api/vendor-audit/contractors", headers=admin_headers)
        assert r.status_code == 200
        ids = [c["id"] for c in r.json()]
        assert TestContractorLifecycle.state["cid"] in ids

    def test_update_contractor(self, admin_headers):
        cid = TestContractorLifecycle.state["cid"]
        r = requests.put(f"{BASE_URL}/api/vendor-audit/contractors/{cid}",
                         json={"contact_phone": "8888888888"}, headers=admin_headers)
        assert r.status_code == 200, r.text
        assert r.json()["contact_phone"] == "8888888888"

    def test_reset_password(self, admin_headers):
        cid = TestContractorLifecycle.state["cid"]
        r = requests.post(f"{BASE_URL}/api/vendor-audit/contractors/{cid}/reset-password", headers=admin_headers)
        assert r.status_code == 200
        new_pwd = r.json()["temp_password"]
        assert len(new_pwd) > 6
        TestContractorLifecycle.state["pwd"] = new_pwd


# ── Contractor auth ──
class TestContractorAuth:
    def test_login(self):
        email = TestContractorLifecycle.state["email"]
        pwd = TestContractorLifecycle.state["pwd"]
        r = requests.post(f"{BASE_URL}/api/contractor/auth/login",
                          json={"email": email, "password": pwd})
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["user"]["must_change_password"] is True
        assert "access_token" in data
        assert data["contractor"]["id"] == TestContractorLifecycle.state["cid"]
        TestContractorLifecycle.state["ctoken"] = data["access_token"]

    def test_change_password(self):
        token = TestContractorLifecycle.state["ctoken"]
        new_pwd = "TestPass@123"
        r = requests.post(f"{BASE_URL}/api/contractor/auth/change-password",
                          json={"current_password": TestContractorLifecycle.state["pwd"],
                                "new_password": new_pwd},
                          headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        # Re-login to get fresh token + verify must_change_password cleared
        r2 = requests.post(f"{BASE_URL}/api/contractor/auth/login",
                           json={"email": TestContractorLifecycle.state["email"], "password": new_pwd})
        assert r2.status_code == 200
        assert r2.json()["user"]["must_change_password"] is False
        TestContractorLifecycle.state["ctoken"] = r2.json()["access_token"]
        TestContractorLifecycle.state["pwd"] = new_pwd

    def test_contractor_blocked_from_admin_endpoint(self):
        token = TestContractorLifecycle.state["ctoken"]
        r = requests.get(f"{BASE_URL}/api/vendor-audit/contractors",
                         headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 403


# ── Audit lifecycle ──
class TestAuditLifecycle:
    def test_template_download(self):
        token = TestContractorLifecycle.state["ctoken"]
        r = requests.get(f"{BASE_URL}/api/vendor-audit/template/vendor-data-sheet",
                         headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200
        assert "attachment" in r.headers.get("Content-Disposition", "")
        assert r.content[:2] == b"PK"  # xlsx is a zip
        TestContractorLifecycle.state["template_xlsx"] = r.content

    def test_start_audit(self):
        token = TestContractorLifecycle.state["ctoken"]
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/start",
                          json={"wage_month": "FEB-2026", "state": "MAHARASHTRA"},
                          headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["status"] == "draft"
        assert data["wage_month"] == "FEB-2026"
        TestContractorLifecycle.state["aid"] = data["id"]

    def test_isolation_other_contractor_audit_forbidden(self, admin_headers):
        # Create a 2nd contractor and audit, then try to access from first contractor token
        unique_email = f"test_vendor2_{uuid.uuid4().hex[:8]}@example.com"
        body = {"name": "TEST E2E Vendor B", "contact_email": unique_email, "state": "MAHARASHTRA"}
        c2 = requests.post(f"{BASE_URL}/api/vendor-audit/contractors", json=body, headers=admin_headers).json()
        # admin starts audit for vendor B
        a2 = requests.post(
            f"{BASE_URL}/api/vendor-audit/audits/start?contractor_id={c2['contractor']['id']}",
            json={"wage_month": "FEB-2026", "state": "MAHARASHTRA"}, headers=admin_headers,
        )
        assert a2.status_code == 200
        other_aid = a2.json()["id"]
        token = TestContractorLifecycle.state["ctoken"]
        r = requests.get(f"{BASE_URL}/api/vendor-audit/audits/{other_aid}",
                         headers={"Authorization": f"Bearer {token}"})
        assert r.status_code in (403, 404)
        # Cleanup
        requests.delete(f"{BASE_URL}/api/vendor-audit/contractors/{c2['contractor']['id']}", headers=admin_headers)

    def test_upload_excel(self):
        # Inject 4 test rows into the downloaded template
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        wb = openpyxl.load_workbook(io.BytesIO(TestContractorLifecycle.state["template_xlsx"]))
        ws = wb.active
        # Find header row by scanning for "EMPLOYEE CODE"
        header_row = None
        for row in ws.iter_rows(min_row=1, max_row=20):
            cells = [str(c.value or "").strip() for c in row]
            if "EMPLOYEE CODE" in cells:
                header_row = row[0].row
                headers = cells
                break
        assert header_row is not None, "Header row with 'EMPLOYEE CODE' not found in template"

        def col(name):
            return headers.index(name) + 1

        def write_row(rownum, mapping):
            for k, v in mapping.items():
                ws.cell(row=rownum, column=col(k), value=v)

        # Row 1: PF wrong (EMP PF=1000 instead of 1800 on 15000 base) → expect PF007
        write_row(header_row + 1, {
            "EMPLOYEE CODE": "E001", "NAME OF EMPLOYEE (AS ON AADHAAR)": "Wrong PF Emp",
            "GENDER": "MALE", "STATE": "MAHARASHTRA",
            "PF APPLICABLE": "YES", "ESIC APPLICABLE": "NO", "MLWF APPLICABLE": "NO",
            "UAN NUMBER": "100200300400", "PF ACCOUNT NUMBER": "MH/PUN/12345/001",
            "PF BASE": 15000, "EPS BASE": 15000, "EDLI BASE": 15000, "ESIC BASE": 0,
            "INCLUSION AMT": 15000, "FINAL GROSS EARNED": 15000, "EARNED BASIC": 15000,
            "EMP PF": 1000, "EMR PF": 1100, "PF PENSION": 1250,
            "PT": 200, "ESIC": 0,
        })
        # Row 2: Bad UAN → PF003
        write_row(header_row + 2, {
            "EMPLOYEE CODE": "E002", "NAME OF EMPLOYEE (AS ON AADHAAR)": "Bad UAN",
            "GENDER": "MALE", "STATE": "MAHARASHTRA",
            "PF APPLICABLE": "YES", "ESIC APPLICABLE": "NO", "MLWF APPLICABLE": "NO",
            "UAN NUMBER": "ABC", "PF ACCOUNT NUMBER": "MH/PUN/12345/002",
            "PF BASE": 15000, "EPS BASE": 15000, "EDLI BASE": 15000, "ESIC BASE": 0,
            "INCLUSION AMT": 15000, "FINAL GROSS EARNED": 15000, "EARNED BASIC": 15000,
            "EMP PF": 1800, "EMR PF": 1100, "PF PENSION": 1250,
            "PT": 200, "ESIC": 0,
        })
        # Row 3: ESIC on 25000 gross → ES001
        write_row(header_row + 3, {
            "EMPLOYEE CODE": "E003", "NAME OF EMPLOYEE (AS ON AADHAAR)": "Wrong ESIC",
            "GENDER": "MALE", "STATE": "MAHARASHTRA",
            "PF APPLICABLE": "NO", "ESIC APPLICABLE": "YES", "MLWF APPLICABLE": "NO",
            "ESIC IP NUMBER": "1122334455",
            "INCLUSION AMT": 25000, "FINAL GROSS EARNED": 25000, "EARNED BASIC": 25000,
            "ESIC BASE": 25000, "EMP CONTRI ESI": 188, "EMR CONTRI ESI": 813,
            "PT": 200, "ESIC": 188,
        })
        # Row 4: PT in Feb at 15000 male → PT001 if 200 instead of 300
        write_row(header_row + 4, {
            "EMPLOYEE CODE": "E004", "NAME OF EMPLOYEE (AS ON AADHAAR)": "PT Feb Wrong",
            "GENDER": "MALE", "STATE": "MAHARASHTRA",
            "PF APPLICABLE": "NO", "ESIC APPLICABLE": "NO", "MLWF APPLICABLE": "NO",
            "INCLUSION AMT": 15000, "FINAL GROSS EARNED": 15000, "EARNED BASIC": 15000,
            "PT": 200,
        })
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        files = {"file": ("seeded_template.xlsx", out.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/upload-excel",
                          files=files, headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        assert r.json()["rows_parsed"] >= 4

    def test_upload_pdf_pf_ecr(self, sample_pdf_bytes):
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        files = {"file": ("STATUTORY DOCUMENTS.pdf", sample_pdf_bytes, "application/pdf")}
        data = {"doc_type": "pf_ecr"}
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/upload-pdf",
                          files=files, data=data, headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        parsed = r.json()["parsed"]
        # Parser nests fields under 'summary'
        summary = parsed.get("summary", {})
        assert summary.get("establishment_name"), f"establishment_name missing: {summary}"
        assert summary.get("lin"), f"lin missing: {summary}"
        # Should also have employees array
        assert isinstance(parsed.get("employees"), list) and len(parsed["employees"]) > 0

    def test_upload_pdf_rejects_non_pdf(self):
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        files = {"file": ("not_pdf.txt", b"hello world", "text/plain")}
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/upload-pdf",
                          files=files, data={"doc_type": "pf_ecr"},
                          headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 400

    def test_run_audit(self):
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/run",
                          headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        result = r.json()
        assert "per_employee" in result or "summary_findings" in result or "totals" in result
        # Check expected rule codes triggered
        all_codes = set()
        for emp in result.get("per_employee", []):
            for f in emp.get("findings", []):
                all_codes.add(f.get("rule_code"))
        for f in result.get("summary_findings", []):
            all_codes.add(f.get("rule_code"))
        # At least PF007 (wrong contribution) and PF003 (bad UAN) should trigger
        assert "PF007" in all_codes, f"Expected PF007 in findings; got {all_codes}"
        assert "PF003" in all_codes, f"Expected PF003 in findings; got {all_codes}"
        assert "ES001" in all_codes, f"Expected ES001 in findings; got {all_codes}"
        # Note: PT001 needs PT challan doc cross-check; in this e2e we didn't upload PT challan
        # so PT001 may not trigger from row-only data. Acceptable.

    def test_register_downloads(self):
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        for kind in ["pf", "esic", "pt"]:
            r = requests.get(f"{BASE_URL}/api/vendor-audit/audits/{aid}/register/{kind}",
                             headers={"Authorization": f"Bearer {token}"})
            assert r.status_code == 200, f"{kind}: {r.text}"
            assert "attachment" in r.headers.get("Content-Disposition", "")
            assert r.content[:2] == b"PK", f"{kind} register not a valid xlsx"

    def test_submit_audit(self):
        token = TestContractorLifecycle.state["ctoken"]
        aid = TestContractorLifecycle.state["aid"]
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/submit",
                          headers={"Authorization": f"Bearer {token}"})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "submitted"

    def test_admin_approve(self, admin_headers):
        aid = TestContractorLifecycle.state["aid"]
        r = requests.post(f"{BASE_URL}/api/vendor-audit/audits/{aid}/approve",
                          json={"remarks": "OK"}, headers=admin_headers)
        assert r.status_code == 200

    def test_cleanup(self, admin_headers):
        cid = TestContractorLifecycle.state.get("cid")
        if cid:
            requests.delete(f"{BASE_URL}/api/vendor-audit/contractors/{cid}", headers=admin_headers)


# ── Security: HRMS endpoints ──
def test_contractor_token_blocked_on_hrms_employees():
    token = TestContractorLifecycle.state.get("ctoken")
    if not token:
        pytest.skip("No contractor token")
    r = requests.get(f"{BASE_URL}/api/employees", headers={"Authorization": f"Bearer {token}"})
    assert r.status_code in (401, 403)

"""Deterministic content extractors for template files.

These run BEFORE the LLM call so we can:
1. Send rich structured context (headers, cells, merged ranges) to the LLM
2. Avoid sending raw Excel bytes to the LLM (Gemini can't natively read .xlsx)
3. Give the LLM an already-parsed view to reason about
"""
from __future__ import annotations
from typing import Dict, Any, List, Optional
import io
import openpyxl
import pdfplumber
from docx import Document


def extract_excel(content: bytes, max_rows: int = 50, max_cols: int = 40) -> Dict[str, Any]:
    """Extract a structured view of an .xlsx template — top rows + merged ranges + sample data."""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    sheets: List[Dict[str, Any]] = []
    for ws in wb.worksheets:
        max_r = min(ws.max_row or 0, max_rows)
        max_c = min(ws.max_column or 0, max_cols)
        rows: List[List[Optional[str]]] = []
        for r in range(1, max_r + 1):
            row = []
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value
                if val is None:
                    row.append(None)
                elif hasattr(cell, "data_type") and cell.data_type == "f":
                    # formula
                    row.append(f"={cell.value.lstrip('=')}" if str(cell.value).startswith("=") else f"={cell.value}")
                else:
                    s = str(val).strip()
                    row.append(s if s else None)
            rows.append(row)
        merged = [str(r) for r in ws.merged_cells.ranges]
        sheets.append({
            "name": ws.title,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "preview_rows": rows,
            "merged_ranges": merged,
        })
    return {"kind": "xlsx", "sheets": sheets}


def extract_pdf(content: bytes, max_pages: int = 5) -> Dict[str, Any]:
    """Extract text + tables from a PDF template."""
    text_pages: List[str] = []
    tables: List[List[List[str]]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages[:max_pages]:
            try:
                txt = page.extract_text() or ""
            except Exception:
                txt = ""
            text_pages.append(txt)
            try:
                page_tables = page.extract_tables() or []
                for t in page_tables:
                    tables.append([[(c or "").strip() for c in row] for row in t])
            except Exception:
                pass
    return {"kind": "pdf", "page_count": len(text_pages), "text_pages": text_pages, "tables": tables}


def extract_docx(content: bytes) -> Dict[str, Any]:
    """Extract paragraphs + tables from a .docx template."""
    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    tables: List[List[List[str]]] = []
    for t in doc.tables:
        tbl_rows = []
        for row in t.rows:
            tbl_rows.append([cell.text.strip() for cell in row.cells])
        tables.append(tbl_rows)
    return {"kind": "docx", "paragraphs": paragraphs, "tables": tables}


def detect_kind(file_name: str) -> str:
    n = (file_name or "").lower()
    if n.endswith((".xlsx", ".xls")):
        return "xlsx"
    if n.endswith(".pdf"):
        return "pdf"
    if n.endswith(".docx"):
        return "docx"
    raise ValueError("Unsupported file type. Please upload .xlsx, .pdf, or .docx")


def extract_any(content: bytes, file_name: str) -> Dict[str, Any]:
    kind = detect_kind(file_name)
    if kind == "xlsx":
        return extract_excel(content)
    if kind == "pdf":
        return extract_pdf(content)
    return extract_docx(content)

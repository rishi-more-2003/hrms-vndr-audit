"""Parsers for the vendor-uploaded monthly Excel + statutory PDFs."""
from __future__ import annotations
import io
import re
from typing import Dict, Any, List, Tuple
import openpyxl
import pdfplumber

from .models import VENDOR_SHEET_COLUMNS


# ── Excel parser ──
def parse_vendor_excel(content: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    """Return (rows, warnings). Rows are dicts keyed by canonical column name."""
    warnings: List[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise ValueError(f"Excel file is corrupt or not a valid .xlsx: {e}")
    ws = wb.active
    # Find header row — the row that contains "EMPLOYEE CODE"
    header_row = None
    for ridx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any("EMPLOYEE CODE" == c.upper() for c in cells):
            header_row = ridx
            headers = cells
            break
    if header_row is None:
        raise ValueError("Could not find header row (looking for 'EMPLOYEE CODE' column)")
    # Normalize headers -> match against canonical list (case-insensitive).
    canonical_lookup = {c.upper().strip(): c for c in VENDOR_SHEET_COLUMNS}
    col_map: Dict[int, str] = {}
    for idx, h in enumerate(headers):
        if not h:
            continue
        key = h.upper().strip()
        # Heuristic for legacy header variants
        if key == "DATE OF JOINIG":  # common typo in user's template
            key = "DATE OF JOINING"
        if key in canonical_lookup:
            col_map[idx] = canonical_lookup[key]
        else:
            # Unknown header — keep as-is in case audit rules want raw access
            col_map[idx] = h.strip()
    rows: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
        # Skip instruction/notes rows (rows where EMPLOYEE CODE is None/empty)
        rec: Dict[str, Any] = {}
        for idx, val in enumerate(row):
            key = col_map.get(idx)
            if not key:
                continue
            rec[key] = val
        # require at least employee_code + name
        if not rec.get("EMPLOYEE CODE") or not rec.get("NAME OF EMPLOYEE (AS ON AADHAAR)"):
            continue
        rows.append(rec)
    if not rows:
        warnings.append("No employee rows detected — check that header row has 'EMPLOYEE CODE' and data rows below it.")
    return rows, warnings


def build_template_xlsx() -> bytes:
    """Generate the vendor data collection Excel template with instructions."""
    from .models import SHEET_INSTRUCTIONS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PAYROLL"
    # Instructions on top (row 1..N)
    ws.cell(row=1, column=1, value="INSTRUCTIONS — read before filling").font = openpyxl.styles.Font(bold=True, size=12)
    for i, line in enumerate(SHEET_INSTRUCTIONS, start=2):
        ws.cell(row=i, column=1, value=line)
    header_row = len(SHEET_INSTRUCTIONS) + 3
    for cidx, col in enumerate(VENDOR_SHEET_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=cidx, value=col)
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="2A2624")
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[header_row].height = 40
    # Freeze top + header
    ws.freeze_panes = ws.cell(row=header_row + 1, column=4)
    # Widths — skinny by default
    for cidx in range(1, len(VENDOR_SHEET_COLUMNS) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cidx)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF parsers ──
def _extract_text_strict(content: bytes) -> str:
    """Extract text from PDF. Reject scanned PDFs (no real text)."""
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            chunks = []
            for page in pdf.pages:
                t = page.extract_text() or ""
                chunks.append(t)
            full = "\n".join(chunks).strip()
    except Exception as e:
        raise ValueError(f"PDF is corrupt or unreadable: {e}")
    if len(full) < 50:
        raise ValueError("This PDF has no extractable text — looks like a scanned/printed copy. Please upload the original PDF downloaded from the govt portal.")
    return full


def _extract_tables(content: bytes) -> List[List[List[str]]]:
    tables: List[List[List[str]]] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            for t in (page.extract_tables() or []):
                tables.append([[(c or "").strip() for c in row] for row in t])
    return tables


def _num(s: str) -> float:
    if s is None:
        return 0.0
    s = str(s).replace(",", "").replace("₹", "").replace("Rs.", "").replace("INR", "").strip()
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        m = re.search(r"-?\d+(?:\.\d+)?", s)
        return float(m.group()) if m else 0.0


def parse_pf_ecr(content: bytes) -> Dict[str, Any]:
    """PF ECR contains establishment details + per-employee contribution rows."""
    text = _extract_text_strict(content)
    tables = _extract_tables(content)
    result: Dict[str, Any] = {"doc_type": "pf_ecr", "summary": {}, "employees": []}
    def find(pattern: str, flags=re.IGNORECASE) -> str:
        m = re.search(pattern, text, flags)
        return m.group(1).strip() if m else ""
    result["summary"]["trrn"] = find(r"TRRN\s*(?:No\.?|Number)?\s*[:\-]?\s*([0-9]{10,})")
    result["summary"]["establishment_name"] = find(r"(?:Name of Establishment|Establishment Name)\s*[:\-]?\s*([A-Z0-9 .&,\-/]+)")
    result["summary"]["establishment_id"] = find(r"(?:Establishment ID|Establishment Code|Estt\.?\s*Code)\s*[:\-]?\s*([A-Z0-9/\-]+)")
    result["summary"]["lin"] = find(r"LIN\s*[:\-]?\s*([0-9]+)")
    result["summary"]["contribution_rate_pct"] = _num(find(r"Contribution Rate\s*\(?%?\)?\s*[:\-]?\s*([\d\.]+)"))
    result["summary"]["total_members"] = int(_num(find(r"Total Members\s*[:\-]?\s*([\d,]+)")))
    result["summary"]["wage_month"] = find(r"(?:Wage Month|For the Month of)\s*[:\-]?\s*([A-Z]{3,}[\s\-,]*\d{4})")
    result["summary"]["total_epf_remitted"] = _num(find(r"Total EPF Contribution Remitted\s*[:\-]?\s*([\d,\.]+)"))
    result["summary"]["total_eps_remitted"] = _num(find(r"Total EPS Contribution Remitted\s*[:\-]?\s*([\d,\.]+)"))
    result["summary"]["total_epf_eps_diff_remitted"] = _num(find(r"Total EPF\s*-\s*EPS Contribution Remitted\s*[:\-]?\s*([\d,\.]+)"))
    # Employee-level rows — look for tables with UAN column
    for tbl in tables:
        if not tbl:
            continue
        header = [c.upper() for c in tbl[0]]
        if any("UAN" in c for c in header):
            uan_idx = next((i for i, c in enumerate(header) if "UAN" in c and "RETURN" not in c), None)
            name_idx = next((i for i, c in enumerate(header) if "NAME" in c), None)
            gross_idx = next((i for i, c in enumerate(header) if "GROSS" in c and "EPF" in c), None)
            epf_idx = next((i for i, c in enumerate(header) if c.strip() in ("EPF", "EPF WAGES") or "EPF WAG" in c), None)
            eps_idx = next((i for i, c in enumerate(header) if c.strip() == "EPS" or "EPS WAG" in c), None)
            edli_idx = next((i for i, c in enumerate(header) if "EDLI" in c), None)
            ee_idx = next((i for i, c in enumerate(header) if c.strip() in ("EE", "EE CONTRIB") or "EMPLOYEE" in c and "PF" in c), None)
            er_idx = next((i for i, c in enumerate(header) if c.strip() == "ER" or ("EMPLOYER" in c and "PF" in c and "PEN" not in c)), None)
            eps_contrib_idx = next((i for i, c in enumerate(header) if "EPS" in c and ("CONTRIB" in c or "PEN" in c)), None)
            ncp_idx = next((i for i, c in enumerate(header) if "NCP" in c), None)
            for r in tbl[1:]:
                if not r or (uan_idx is not None and not r[uan_idx]):
                    continue
                result["employees"].append({
                    "uan": r[uan_idx] if uan_idx is not None and uan_idx < len(r) else "",
                    "name": r[name_idx] if name_idx is not None and name_idx < len(r) else "",
                    "gross_epf": _num(r[gross_idx]) if gross_idx is not None and gross_idx < len(r) else 0.0,
                    "epf_wages": _num(r[epf_idx]) if epf_idx is not None and epf_idx < len(r) else 0.0,
                    "eps_wages": _num(r[eps_idx]) if eps_idx is not None and eps_idx < len(r) else 0.0,
                    "edli_wages": _num(r[edli_idx]) if edli_idx is not None and edli_idx < len(r) else 0.0,
                    "ee_contrib": _num(r[ee_idx]) if ee_idx is not None and ee_idx < len(r) else 0.0,
                    "er_contrib": _num(r[er_idx]) if er_idx is not None and er_idx < len(r) else 0.0,
                    "eps_contrib": _num(r[eps_contrib_idx]) if eps_contrib_idx is not None and eps_contrib_idx < len(r) else 0.0,
                    "ncp_days": int(_num(r[ncp_idx])) if ncp_idx is not None and ncp_idx < len(r) else 0,
                })
    return result


def parse_pf_challan(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    return {
        "doc_type": "pf_challan",
        "summary": {
            "trrn": find(r"TRRN\s*(?:No\.?|Number)?\s*[:\-]?\s*([0-9]{10,})"),
            "establishment_code": find(r"(?:Establishment\s*Details\s*-\s*Code|Establishment\s*Code)\s*[:\-]?\s*([A-Z0-9/\-]+)"),
            "establishment_name": find(r"(?:Establishment\s*Details\s*-\s*Name|Establishment\s*Name)\s*[:\-]?\s*([A-Z0-9 .&,\-/]+)"),
            "establishment_address": find(r"(?:Establishment\s*Details\s*-\s*Address|Address)\s*[:\-]?\s*([A-Za-z0-9 .,\-/#]+)"),
            "total_subscribers": int(_num(find(r"Total Subscribers\s*[:\-]?\s*([\d,]+)"))),
            "payment_date": find(r"Payment\s*Details\s*-?\s*Date\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "ac01_emp": _num(find(r"A/?C\.?\s*0?1\s+.*?Employee.*?Share\s*[:\-]?\s*([\d,\.]+)", re.IGNORECASE | re.DOTALL)),
            "ac01_er": _num(find(r"A/?C\.?\s*0?1\s+.*?Employer.*?Share\s*[:\-]?\s*([\d,\.]+)", re.IGNORECASE | re.DOTALL)),
            "ac02_admin": _num(find(r"A/?C\.?\s*0?2\s+.*?Admin.*?Charges\s*[:\-]?\s*([\d,\.]+)", re.IGNORECASE | re.DOTALL)),
            "ac10_er": _num(find(r"A/?C\.?\s*10\s+.*?Employer.*?Share\s*[:\-]?\s*([\d,\.]+)", re.IGNORECASE | re.DOTALL)),
            "ac21_er": _num(find(r"A/?C\.?\s*21\s+.*?Employer.*?Share\s*[:\-]?\s*([\d,\.]+)", re.IGNORECASE | re.DOTALL)),
        },
    }


def parse_pf_paid_challan(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    return {
        "doc_type": "pf_paid_challan",
        "summary": {
            "trrn": find(r"TRRN\s*(?:No\.?|Number)?\s*[:\-]?\s*([0-9]{10,})"),
            "payment_date": find(r"(?:Paid\s*(?:On|Date)|Payment\s*Date)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "amount_paid": _num(find(r"(?:Amount\s*Paid|Total\s*Paid|Amount)\s*[:\-]?\s*([\d,\.]+)")),
            "status": find(r"(?:Status|Payment Status)\s*[:\-]?\s*([A-Z]+)"),
            "bank_cin": find(r"(?:Bank\s*CIN|CIN)\s*[:\-]?\s*([A-Z0-9]+)"),
        },
    }


def parse_esic_contribution_history(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    tables = _extract_tables(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    result = {
        "doc_type": "esic_contribution_history",
        "summary": {
            "employer_code": find(r"(?:Employer\s*Code|Employer\s*Code\s*No\.?)\s*[:\-]?\s*([0-9]{10,})"),
            "employer_name": find(r"(?:Employer\s*Name|Name\s*of\s*the\s*Employer)\s*[:\-]?\s*([A-Z0-9 .&,\-/]+)"),
            "wage_month": find(r"(?:Contribution\s*Period|For\s*the\s*Month)\s*[:\-]?\s*([A-Z]{3,}[\s\-]*\d{2,4})"),
            "total_ip": int(_num(find(r"Total\s*(?:No\s*of)?\s*IP\s*[:\-]?\s*([\d,]+)"))),
            "total_wages": _num(find(r"Total\s*(?:IP)?\s*Wages\s*[:\-]?\s*([\d,\.]+)")),
            "total_contribution": _num(find(r"Total\s*Contribution\s*[:\-]?\s*([\d,\.]+)")),
        },
        "employees": [],
    }
    for tbl in tables:
        if not tbl:
            continue
        hdr = [c.upper() for c in tbl[0]]
        if any("IP NUMBER" in c or "INSURANCE NO" in c or c.strip() == "IP" for c in hdr):
            ip_idx = next((i for i, c in enumerate(hdr) if "IP" in c and ("NUMBER" in c or "NO" in c)), None)
            name_idx = next((i for i, c in enumerate(hdr) if "NAME" in c), None)
            days_idx = next((i for i, c in enumerate(hdr) if "DAY" in c), None)
            wages_idx = next((i for i, c in enumerate(hdr) if "WAGES" in c or "WAGE" in c), None)
            contrib_idx = next((i for i, c in enumerate(hdr) if "CONTRIB" in c), None)
            for r in tbl[1:]:
                if not r or (ip_idx is not None and not r[ip_idx]):
                    continue
                result["employees"].append({
                    "ip_number": r[ip_idx] if ip_idx is not None and ip_idx < len(r) else "",
                    "name": r[name_idx] if name_idx is not None and name_idx < len(r) else "",
                    "days": int(_num(r[days_idx])) if days_idx is not None and days_idx < len(r) else 0,
                    "wages": _num(r[wages_idx]) if wages_idx is not None and wages_idx < len(r) else 0.0,
                    "contribution": _num(r[contrib_idx]) if contrib_idx is not None and contrib_idx < len(r) else 0.0,
                })
    return result


def parse_esic_paid_challan(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    return {
        "doc_type": "esic_paid_challan",
        "summary": {
            "challan_no": find(r"Challan\s*(?:No|Number)\s*[:\-]?\s*([A-Z0-9/]+)"),
            "employer_code": find(r"(?:Employer\s*Code|Code No)\s*[:\-]?\s*([0-9]+)"),
            "payment_date": find(r"(?:Paid\s*(?:On|Date)|Payment\s*Date|Date\s*of\s*Payment)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "amount_paid": _num(find(r"(?:Amount\s*Paid|Total\s*Contribution|Total\s*Amount)\s*[:\-]?\s*([\d,\.]+)")),
            "status": find(r"(?:Status|Payment Status)\s*[:\-]?\s*([A-Z]+)"),
        },
    }


def parse_pt_paid_challan(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    return {
        "doc_type": "pt_paid_challan",
        "summary": {
            "grn": find(r"(?:GRN|CIN|Challan No)\s*[:\-]?\s*([A-Z0-9]+)"),
            "tin": find(r"(?:TIN|PT\s*Registration\s*No|Registration\s*No)\s*[:\-]?\s*([A-Z0-9/]+)"),
            "mvat_gstn": find(r"(?:MVAT|GSTIN|GSTN)\s*[:\-]?\s*([A-Z0-9]+)"),
            "employer_name": find(r"(?:Name of the Employer|Employer Name)\s*[:\-]?\s*([A-Z0-9 .&,\-/]+)"),
            "payment_date": find(r"(?:Payment\s*Date|Date\s*of\s*Payment|Paid\s*On)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "amount_paid": _num(find(r"(?:Amount\s*Paid|Total\s*Amount|Tax\s*Amount)\s*[:\-]?\s*([\d,\.]+)")),
            "period_from": find(r"(?:Period.*?From)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "period_to": find(r"(?:Period.*?To|To\s*Date)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
        },
    }


def parse_pt_return(content: bytes) -> Dict[str, Any]:
    text = _extract_text_strict(content)
    def find(p, f=re.IGNORECASE):
        m = re.search(p, text, f); return m.group(1).strip() if m else ""
    return {
        "doc_type": "pt_return",
        "summary": {
            "tin": find(r"(?:TIN|PT\s*Reg)\s*[:\-]?\s*([A-Z0-9/]+)"),
            "employer_name": find(r"(?:Name\s*of\s*(?:the\s*)?Employer)\s*[:\-]?\s*([A-Z0-9 .&,\-/]+)"),
            "return_type": find(r"(?:Type\s*of\s*Return)\s*[:\-]?\s*([A-Z ]+)"),
            "period_from": find(r"(?:Period.*?From)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "period_to": find(r"(?:Period.*?To)\s*[:\-]?\s*(\d{1,2}[/\-]\d{1,2}[/\-]\d{2,4})"),
            "total_tax": _num(find(r"(?:Total\s*Tax|Tax\s*Amount)\s*[:\-]?\s*([\d,\.]+)")),
        },
    }


DOC_PARSERS = {
    "pf_ecr": parse_pf_ecr,
    "pf_challan": parse_pf_challan,
    "pf_paid_challan": parse_pf_paid_challan,
    "esic_contribution_history": parse_esic_contribution_history,
    "esic_paid_challan": parse_esic_paid_challan,
    "pt_paid_challan": parse_pt_paid_challan,
    "pt_return": parse_pt_return,
}

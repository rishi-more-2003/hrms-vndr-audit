"""PF / ESIC / PT summary register generators as Excel exports."""
from __future__ import annotations
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from typing import List, Dict, Any


def _styled_header(ws, row, cols):
    for c, v in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2A2624")
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 30


def _n(v):
    try: return float(str(v).replace(",", "").strip() or 0)
    except: return 0


def build_pf_register(rows: List[Dict[str, Any]], vendor: Dict[str, Any], month: str) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "PF Register"
    ws.cell(row=1, column=1, value=f"PF Register — {vendor.get('name','')} — {month}").font = Font(bold=True, size=14)
    headers = ["Sr No", "Employee Code", "Name", "UAN", "PF A/C No", "PF Base", "EPS Base", "EDLI Base",
               "Employee PF (12%)", "Employer PF (3.67%)", "EPS (8.33%)", "EDLI (0.5%)", "PF Admin (0.5%)", "Total"]
    _styled_header(ws, 3, headers)
    for i, r in enumerate(rows, start=1):
        pf_base = _n(r.get("PF BASE"))
        eps_base = _n(r.get("EPS BASE"))
        emp_pf = _n(r.get("EMP PF"))
        er_pf = _n(r.get("EMR PF"))
        eps = _n(r.get("PF PENSION"))
        edli = round(_n(r.get("EDLI BASE")) * 0.005)
        admin = round(pf_base * 0.005)
        ws.append([
            i, r.get("EMPLOYEE CODE"), r.get("NAME OF EMPLOYEE (AS ON AADHAAR)"),
            r.get("UAN NUMBER"), r.get("PF ACCOUNT NUMBER"),
            pf_base, eps_base, _n(r.get("EDLI BASE")),
            emp_pf, er_pf, eps, edli, admin,
            round(emp_pf + er_pf + eps + edli + admin),
        ])
    for cidx in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cidx)].width = 15
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_esic_register(rows: List[Dict[str, Any]], vendor: Dict[str, Any], month: str) -> bytes:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "ESIC Register"
    ws.cell(row=1, column=1, value=f"ESIC Register — {vendor.get('name','')} — {month}").font = Font(bold=True, size=14)
    headers = ["Sr No", "Employee Code", "Name", "IP Number", "ESIC Base", "Days",
               "Employee ESIC (0.75%)", "Employer ESIC (3.25%)", "Total"]
    _styled_header(ws, 3, headers)
    for i, r in enumerate(rows, start=1):
        base = _n(r.get("ESIC BASE"))
        emp_esic = _n(r.get("ESIC"))
        er_esic = round(base * 0.0325)
        ws.append([
            i, r.get("EMPLOYEE CODE"), r.get("NAME OF EMPLOYEE (AS ON AADHAAR)"),
            r.get("ESIC IP NUMBER"), base, _n(r.get("PAID DAYS")),
            emp_esic, er_esic, round(emp_esic + er_esic),
        ])
    for cidx in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cidx)].width = 16
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def build_pt_register(rows: List[Dict[str, Any]], vendor: Dict[str, Any], month: str) -> bytes:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "PT Register"
    ws.cell(row=1, column=1, value=f"Professional Tax Register — {vendor.get('name','')} — {month}").font = Font(bold=True, size=14)
    headers = ["Sr No", "Employee Code", "Name", "Gender", "State", "Final Gross Earned", "PT Slab", "PT Amount"]
    _styled_header(ws, 3, headers)
    from .rules import _pt_mh_expected
    for i, r in enumerate(rows, start=1):
        g = _n(r.get("FINAL GROSS EARNED"))
        gender = (r.get("GENDER") or "").upper()
        state = (r.get("STATE") or "").upper()
        expected = _pt_mh_expected(gender, (r.get("WAGE MONTH") or "").upper(), g) if "MAHARASHTRA" in state else None
        slab = "N/A" if expected is None else ("0" if expected == 0 else f"₹{expected}")
        ws.append([
            i, r.get("EMPLOYEE CODE"), r.get("NAME OF EMPLOYEE (AS ON AADHAAR)"),
            gender, state, g, slab, _n(r.get("PT")),
        ])
    for cidx in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(cidx)].width = 17
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
